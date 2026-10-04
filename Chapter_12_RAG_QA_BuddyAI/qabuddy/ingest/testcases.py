"""Test-case spreadsheets -> one document per row (never split)."""

from __future__ import annotations

import csv
from pathlib import Path

from .base import document, iter_files

SUFFIXES = [".csv", ".xlsx", ".xlsm"]

# Canonical metadata we lift out of the row (matched case-insensitively).
# Aliases cover both the Chapter 11 export ("Test Case ID", "Summary") and the
# VWO 500-case export ("Scenario TID", "TestCase Description").
_META_COLUMNS = {
    "tc_id": {"test case id", "tc id", "id", "testcase id", "scenario tid", "scenario id"},
    "summary": {
        "summary",
        "title",
        "test case summary",
        "testcase description",
        "test case description",
    },
    "priority": {"priority"},
    "category": {"category", "module"},
    "scenario_type": {"scenario type"},
    "test_type": {"test type"},
    "execution_status": {"execution status", "status"},
    "environment": {"environment", "env"},
    "precondition": {"precondition", "preconditions"},
    "is_automated": {"is automated", "automated"},
}


def _row_text(row: dict[str, str]) -> str:
    lines = []
    for key, value in row.items():
        if key is None:
            continue
        value = (value or "").strip()
        if value:
            lines.append(f"{key}: {value}")
    return "\n\n".join(lines)


def _meta_lookup(row: dict[str, str]) -> dict[str, str]:
    lowered = {str(k).strip().lower(): (v or "").strip() for k, v in row.items() if k}
    meta: dict[str, str] = {}
    for field, names in _META_COLUMNS.items():
        for name in names:
            if lowered.get(name):
                meta[field] = lowered[name]
                break
    return meta


def _csv_rows(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            yield row


def _xlsx_rows(path: Path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        try:
            header = [str(c).strip() if c is not None else "" for c in next(rows)]
        except StopIteration:
            continue
        for values in rows:
            yield {header[i]: ("" if v is None else str(v)) for i, v in enumerate(values)}


def load(settings, folder: str) -> list[dict]:
    base = settings.source_dir(folder)
    docs: list[dict] = []
    for path in iter_files(base, SUFFIXES):
        rows = _xlsx_rows(path) if path.suffix.lower() in {".xlsx", ".xlsm"} else _csv_rows(path)
        for index, row in enumerate(rows):
            text = _row_text(row)
            if not text.strip():
                continue
            meta = _meta_lookup(row)
            key = meta.get("tc_id") or f"{path.stem}-{index}"
            docs.append(
                document(
                    text,
                    point_key=f"testcases:{path.name}:{key}",
                    source_type="test_case",
                    source_file=path.name,
                    **meta,
                )
            )
    return docs
